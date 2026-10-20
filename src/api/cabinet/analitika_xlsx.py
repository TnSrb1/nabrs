"""Excel export of the “Analitika” report builder (R4).

A pure builder (no DB/network): header + rows + notes → .xlsx bytes.
Mode ``write_only`` (openpyxl) — streaming write, memory does not grow with the number of
rows (the cap of 50 000 is set by the service).

Two sheets with fixed names: ``Data`` (report rows) and ``Napomena``
(the set, the filters, the data period, the UTC moment of export, the share of rows without an amount,
sources). Without “Napomena” the report lies: a third of the tenders have no amount.

Numbers are written as numbers, dates as dates (``date``/``datetime`` + cell format),
so that Excel sorts and sums without “text that looks like a number”.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from src.domain.text_norm import sr_display

DATA_SHEET = "Data"
NOTES_SHEET = "Napomena"

_DATE_FMT = "yyyy-mm-dd"
_MONTH_FMT = "yyyy-mm"
_DATETIME_FMT = "yyyy-mm-dd hh:mm"


class MonthValue(date):
    """The date of the first day of the month, which Excel shows as ``yyyy-mm``."""


#: The first characters with which Excel/LibreOffice start a formula (CSV/formula
#: injection): the captions of customers/winners come from the portal, the filters — from
#: the user. Such a string is written EXPLICITLY as text, not as a formula.
_FORMULA_LEAD = ("=", "+", "-", "@", "\t", "\r")


def _cell(ws, value: Any, lang: str | None = None):
    value = sr_display(value, lang)
    if isinstance(value, Decimal):
        value = float(value)
    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, str) and value.startswith(_FORMULA_LEAD):
        cell.data_type = "s"
    if isinstance(value, datetime):
        cell.number_format = _DATETIME_FMT
    elif isinstance(value, MonthValue):
        cell.number_format = _MONTH_FMT
    elif isinstance(value, date):
        cell.number_format = _DATE_FMT
    return cell


def build_xlsx(
    header: Sequence[str],
    rows: Iterable[Sequence[Any]],
    notes: Iterable[tuple[str, Any]],
    lang: str | None = None,
) -> bytes:
    """Assembles the workbook: sheet ``Data`` (header + rows) and ``Napomena`` (key/value).

    ``lang == "sr"`` — all text cells in Latin script (owner 10.10.2026).
    """
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(DATA_SHEET)
    ws.append([_cell(ws, h, lang) for h in header])
    for r in rows:
        ws.append([_cell(ws, v, lang) for v in r])
    ns = wb.create_sheet(NOTES_SHEET)
    for key, value in notes:
        ns.append([_cell(ns, key, lang), _cell(ns, value, lang)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
