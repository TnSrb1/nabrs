"""Tests of “Analitika” (R4): routes, gate, allowlist, Excel, timeout.

SQL aggregates on a live PG — ``tests/db/test_analitika_report_sql.py``. Here:
the HTTP contract of the routes (the service is mocked), the pure gate, parsing of the input up to the
allowlist, the assembly of the .xlsx (opened by openpyxl) and the timeout path on a mock session.
"""

from __future__ import annotations

import base64
import hashlib
import io
import re
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.exc import DBAPIError

from src.analytics import report_builder as rb
from src.api import deps, security
from src.api.cabinet import analitika_service as svc
from src.api.cabinet import routes_analitika
from src.api.cabinet.analitika_xlsx import build_xlsx
from src.api.cabinet.analitika_service import (
    AnalitikaExportView,
    AnalitikaPageView,
    RowView,
    analytics_gate,
)
from src.api.config import get_cabinet_settings
from src.api.main import app
from tests.api.conftest import FakeSession

SESS = {"uid": "u-1", "tid": "t-1"}
NOW = datetime(2026, 10, 10, 12, 0, tzinfo=timezone.utc)


def _client(*, sess=None):
    app.dependency_overrides[deps.get_session_factory] = lambda: FakeSession
    app.dependency_overrides[deps.get_optional_session] = lambda: sess
    client = TestClient(app, follow_redirects=False)
    if sess:
        settings = get_cabinet_settings()
        client.cookies.set(
            "tb_session",
            security.sign_session(
                SESS["uid"], SESS["tid"], settings.cabinet_secret_key, locale="sr"
            ),
        )
    return client


def teardown_function():
    app.dependency_overrides.clear()


# --------------------------------------------------------------------------- #
#  Gate (new axis: Business/Enterprise or trial; freezing first)
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    "status,tier,expected",
    [
        ("active", "pro", "gate_business"),
        ("active", "lite", "gate_business"),
        ("active", "business", None),
        ("active", "enterprise", None),
        ("trial", "lite", None),  # trial = Business analytics (owner 10.10)
        ("trial", "pro", None),
    ],
)
def test_gate_axis(status, tier, expected):
    assert analytics_gate(status, tier, None, NOW) == expected


def test_gate_frozen_checked_first():
    past = NOW - timedelta(seconds=1)
    assert analytics_gate("trial", "lite", past, NOW) == "frozen"
    assert analytics_gate("active", "business", past, NOW) == "frozen"
    assert analytics_gate("active", "business", NOW + timedelta(days=3), NOW) is None


# --------------------------------------------------------------------------- #
#  Allowlist: the input does not get into SQL
# --------------------------------------------------------------------------- #


def test_parse_spec_rejects_unknown_and_injection():
    spec = rb.parse_spec(
        {
            "subject": "tenders; DROP TABLE x",
            "g1": "customer_name) --",
            "g2": "winner",  # winner — only on decisions
            "m": "winners",  # same
            "cpv": "45'--",
            "region": "RS-00' OR 1=1",
            "amin": "abc",
            "profile": "../../etc",
        }
    )
    assert spec.subject == "tenders"
    assert spec.group1 == "customer"
    assert spec.group2 == ""
    assert spec.measure == "count"
    assert spec.cpv == ("45",)
    assert spec.region == ""
    assert spec.amount_min is None
    assert spec.profile_id == ""


def test_user_values_are_bound_not_inlined():
    spec = rb.parse_spec(
        {"subject": "awards", "g1": "winner", "customer": "x' OR '1'='1", "winner": "Y"}
    )
    sql, params = rb.build_report_sql(spec, limit=50)
    assert "OR '1'='1" not in sql
    assert params["customer_q"] == "x' or '1'='1"
    assert ":customer_q" in sql and ":winner_q" in sql


def test_sql_carries_dataset_predicates():
    """Text guard: the set predicates in SQL (the behavior — in tests/db)."""
    t_sql, _ = rb.build_report_sql(rb.parse_spec({"subject": "tenders"}), limit=5)
    assert "t.published_at >= t.first_seen_at - interval '365 days'" in t_sql
    a_sql, _ = rb.build_report_sql(rb.parse_spec({"subject": "awards"}), limit=5)
    assert "->>'IsReversed', 'false') <> 'true'" in a_sql


def test_profile_filters_use_matching_prefix_canon():
    cpv, regions = rb.profile_filters(
        [
            {"system_code": "eu_cpv", "code": "30000000"},
            {"system_code": "eu_cpv", "code": "45210000"},
            {"system_code": "eu_cpv", "code": "33*"},
            {"system_code": "other", "code": "99"},
        ],
        ["RS-00", "RS-06"],
    )
    assert cpv == ("30", "4521", "33")
    assert regions == ("RS-00", "RS-06")


def test_profile_filters_reads_object_canon_of_prod():
    """The production canon (12/12 live profiles): ``{"include": [...], "exclude": [...]}``.

    Previously the object was iterated by keys → the template’s CPV was not applied at all.
    ``exclude`` is not applied here (see REPORT): only include.
    """
    cpv, _ = rb.profile_filters(
        {"exclude": [], "include": [{"code": "45*", "system_code": "eu_cpv"}]},
        [],
    )
    assert cpv == ("45",)


def test_profile_filters_object_ignores_exclude_for_cpv():
    cpv, _ = rb.profile_filters(
        {
            "include": [{"code": "33*", "system_code": "eu_cpv"}],
            "exclude": [{"code": "3300*", "system_code": "eu_cpv"}],
        },
        [],
    )
    assert cpv == ("33",)


# --------------------------------------------------------------------------- #
#  Routes
# --------------------------------------------------------------------------- #


def test_anon_redirects_to_login():
    c = _client(sess=None)
    assert c.get("/app/analitika").headers["location"] == "/login"
    resp = c.get("/app/analitika.xlsx")
    assert resp.status_code == 303 and resp.headers["location"] == "/login"


def _patch_page(monkeypatch, view):
    async def fake(db, **kwargs):
        return view

    monkeypatch.setattr(routes_analitika.svc, "get_analitika_page", fake)


def test_pro_sees_business_cta(monkeypatch):
    _patch_page(monkeypatch, AnalitikaPageView(status="gate_business"))
    resp = _client(sess=SESS).get("/app/analitika?lang=sr")
    assert resp.status_code == 200
    assert "Izveštaji analitike su deo paketa Business" in resp.text
    assert 'href="/plans"' in resp.text
    assert 'name="subject"' not in resp.text  # no constructor form


def test_business_sees_table_with_totals(monkeypatch):
    rows = [
        RowView("Opština Alfa", "", 2, 1, "100 €", "100 €", "100 €", 1, 0),
        RowView("Bolnica Beta", "", 1, 1, "300 €", "300 €", "300 €", 1, 0),
    ]
    total = RowView("", "", 3, 2, "400 €", "200 €", "200 €", 2, 0)
    _patch_page(
        monkeypatch,
        AnalitikaPageView(
            status="success",
            spec=rb.parse_spec({"subject": "tenders"}),
            heads=svc.sort_heads(rb.parse_spec({"subject": "tenders"})),
            ran=True,
            rows=rows,
            total=total,
            total_groups=2,
            first_date_fmt="10.06.2026",
            last_date_fmt="05.07.2026",
            export_href="/app/analitika.xlsx?subject=tenders",
        ),
    )
    resp = _client(sess=SESS).get("/app/analitika?subject=tenders&lang=sr")
    assert resp.status_code == 200
    assert "Opština Alfa" in resp.text and "Bolnica Beta" in resp.text
    assert "Ukupno" in resp.text
    assert "Iznos je poznat za 2 od 3 zapisa" in resp.text
    assert "10.06.2026 — 05.07.2026" in resp.text
    assert 'href="/app/analitika.xlsx?subject=tenders"' in resp.text
    # Cabinet menu item.
    assert 'href="/app/analitika"' in resp.text


def _fake_result() -> rb.ReportResult:
    row = rb.ReportRow(
        g1_key="2026-06",
        g1_label="2026-06",
        g2_key=None,
        g2_label=None,
        n=2,
        with_amount=1,
        sum_eur=Decimal("100.50"),
        avg_eur=Decimal("100.50"),
        median_eur=100.5,
        distinct_customers=1,
        distinct_winners=0,
    )
    return rb.ReportResult(
        rows=[row],
        total_groups=1,
        totals=rb.ReportTotals(
            n=2,
            with_amount=1,
            sum_eur=Decimal("100.50"),
            avg_eur=Decimal("100.50"),
            median_eur=100.5,
            distinct_customers=1,
            first_date=date(2026, 6, 10),
            last_date=date(2026, 6, 20),
        ),
        coverage=rb.Coverage(
            n=16047,
            with_amount=10591,
            first_date=date(2026, 5, 30),
            last_date=date(2026, 10, 9),
            sources=("jnportal", "ted"),
        ),
    )


def test_xlsx_download_opens_with_data_and_notes(monkeypatch):
    spec = rb.parse_spec({"subject": "tenders", "g1": "month", "cpv": "45"})
    content = svc.build_export(
        spec, _fake_result(), lang="sr", generated_at=NOW, profile_name=""
    )

    async def fake(db, **kwargs):
        return AnalitikaExportView(
            status="success",
            filename=svc.export_filename(spec, NOW),
            content=content,
        )

    monkeypatch.setattr(routes_analitika.svc, "get_analitika_export", fake)
    resp = _client(sess=SESS).get("/app/analitika.xlsx?subject=tenders&g1=month")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert resp.headers["content-disposition"] == (
        'attachment; filename="analitika_tenders_month_20261010.xlsx"'
    )

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Data", "Napomena"]
    data = list(wb["Data"].iter_rows(values_only=True))
    assert data[0][:3] == ("Mesec", "Broj", "Sa iznosom")
    # Month — as a date, numbers — as numbers (not as text).
    assert data[1][0] == datetime(2026, 6, 1)
    assert data[1][1:5] == (2, 1, 100.5, 100.5)
    assert data[2][0] == "Ukupno"
    notes = {r[0]: r[1] for r in wb["Napomena"].iter_rows(values_only=True)}
    assert notes["Skup podataka"] == "Tenderi (objave)"
    assert notes["Filteri"] == "CPV 45*"
    assert notes["Podaci u bazi od"] == datetime(2026, 5, 30)
    assert notes["Udeo zapisa bez iznosa, %"] == 50.0
    assert notes["Izvori"] == "jnportal, ted"
    assert notes["Izvezeno (UTC)"] == datetime(2026, 10, 10, 12, 0)


@pytest.mark.parametrize("status", ["gate_business", "frozen"])
def test_xlsx_gated_redirects_to_plans(monkeypatch, status):
    async def fake(db, **kwargs):
        return AnalitikaExportView(status=status)

    monkeypatch.setattr(routes_analitika.svc, "get_analitika_export", fake)
    resp = _client(sess=SESS).get("/app/analitika.xlsx?subject=tenders")
    assert resp.status_code == 303
    assert resp.headers["location"] == "/plans"


# --------------------------------------------------------------------------- #
#  Timeout: 57014 → rollback and a page with a hint, not 500
# --------------------------------------------------------------------------- #


class _Canceled(Exception):
    sqlstate = "57014"


class _TimeoutSession:
    def __init__(self):
        self.rolled_back = 0
        self.calls: list[str] = []

    async def execute(self, statement, params=None):
        sql = str(statement)
        self.calls.append(sql)
        if "FROM tenants" in sql:
            return SimpleNamespace(
                first=lambda: SimpleNamespace(
                    tier="business", status="active", access_until=None
                )
            )
        if "FROM profiles" in sql:
            return SimpleNamespace(fetchall=lambda: [])
        if "set_config" in sql:
            return SimpleNamespace()
        raise DBAPIError("SELECT", {}, _Canceled())

    async def rollback(self):
        self.rolled_back += 1


@pytest.mark.asyncio
async def test_timeout_rolls_back_and_flags_page():
    db = _TimeoutSession()
    page = await svc.get_analitika_page(
        db, tid="t-1", params={"subject": "tenders"}, lang="sr"
    )
    assert page.status == "success"
    assert page.timed_out is True
    assert db.rolled_back == 1
    # set_tenant — the first action of the transaction.
    assert "app.current_tenant_id" in db.calls[0]


@pytest.mark.parametrize(
    "lang,heading",
    [
        ("sr", "Gotovi izveštaji"),
        ("ru", "Готовые отчёты"),
        ("en", "Ready-made reports"),
    ],
)
def test_form_renders_in_all_languages(monkeypatch, lang, heading):
    _patch_page(
        monkeypatch,
        AnalitikaPageView(status="success", spec=rb.parse_spec({}), presets=[]),
    )
    resp = _client(sess=SESS).get(f"/app/analitika?lang={lang}")
    assert resp.status_code == 200
    assert heading in resp.text
    assert 'name="g1"' in resp.text and 'name="m"' in resp.text


@pytest.mark.parametrize(
    "raw,expected",
    [
        # Output of the CPV widget: rows eu_cpv:CODE; the full code is trimmed to the significant part.
        ("eu_cpv:45210000\neu_cpv:33*", ("4521", "33")),
        (
            "eu_cpv:30000000",
            ("30",),
        ),  # not “3” — otherwise the prefix of all sections 30–39
        ("45, 45;3019", ("45", "3019")),  # duplicates and separators
        ("45'--", ("45",)),
        ("abc", ()),
    ],
)
def test_cpv_widget_input_parsed_to_prefixes(raw, expected):
    assert rb.parse_spec({"cpv": raw}).cpv == expected


def test_cpv_prefixes_capped():
    raw = "\n".join(f"eu_cpv:{i:02d}*" for i in range(10, 60))
    assert len(rb.parse_spec({"cpv": raw}).cpv) <= rb.MAX_CPV_PREFIXES


def test_cpv_filter_is_bound_array_of_prefixes():
    sql, params = rb.build_report_sql(
        rb.parse_spec({"subject": "tenders", "cpv": "eu_cpv:33*\neu_cpv:4521"}),
        limit=5,
    )
    assert params["cpv_like"] == ["33%", "4521%"]
    assert "LIKE ANY(CAST(:cpv_like AS text[]))" in sql
    assert "33%" not in sql


def test_form_uses_cpv_widget_with_versioned_dictionary(monkeypatch):
    spec = rb.parse_spec({"cpv": "33,4521"})
    _patch_page(monkeypatch, AnalitikaPageView(status="success", spec=spec, presets=[]))
    html = _client(sess=SESS).get("/app/analitika?lang=sr").text
    # A widget instead of a text field; the selected codes — line by line in a hidden field.
    assert 'id="cpvIncludeWidget"' in html
    assert 'name="cpv" id="cpvIncludeHidden"' in html
    assert "eu_cpv:33*\neu_cpv:4521*" in html
    assert 'id="an-cpv"' not in html
    # The directory — by the address with ?v=<hash> (#637), not by a bare path.
    m = re.search(r'data-cpv-src="(/static/cpv_codes\.json\?v=[0-9a-f]{10})"', html)
    assert m, "data-cpv-src without a version"
    assert "/js/cpv_widget.js?v=" in html
    # Serbian interface — in Latin script: the label and the hint without Cyrillic.
    block = html.split('id="cpvIncludeWidget"', 1)[1].split("</textarea>", 1)[0]
    assert "CPV kod" in html
    assert not re.search(r"[А-Яа-я]", block)


def test_widget_js_takes_dictionary_url_from_page_not_hardcoded():
    js = (
        Path(__file__).resolve().parents[2] / "src/api/static/js/cpv_widget.js"
    ).read_text(encoding="utf-8")
    assert "getAttribute('data-cpv-src')" in js
    # The bare path is allowed only as a fallback, and in exactly one place.
    assert js.count("/static/cpv_codes.json") == 1
    form = (
        Path(__file__).resolve().parents[2] / "src/api/templates/profile_form.html"
    ).read_text(encoding="utf-8")
    assert "data-cpv-src=\"{{ static_url('cpv_codes.json') }}\"" in form


def test_analitika_widget_sri_matches_file():
    root = Path(__file__).resolve().parents[2]
    digest = hashlib.sha384(
        (root / "src/api/static/js/cpv_widget.js").read_bytes()
    ).digest()
    expected = "sha384-" + base64.b64encode(digest).decode()
    html = (root / "src/api/templates/analitika.html").read_text(encoding="utf-8")
    m = re.search(
        r"static_url\('js/cpv_widget\.js'\) \}\}\" integrity=\"([^\"]+)\"", html
    )
    assert m and m.group(1) == expected


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("1e999999", None),  # otherwise DataError/OverflowError → 500
        ("123456789012345678", None),  # order 17 > 15
        ("1e15", Decimal("1E+15")),
        ("2 500,50", Decimal("2500.50")),
        ("-5", None),
    ],
)
def test_amount_filter_bounded(raw, expected):
    assert rb.parse_spec({"amin": raw}).amount_min == expected


def test_xlsx_formula_injection_written_as_text():
    content = build_xlsx(
        ["=SUM(1)"],
        [['=HYPERLINK("x")', 1], ["+381 11", 2], ["@cmd", 3], ["Obična", 4]],
        [("k", "-2+3")],
    )
    wb = load_workbook(io.BytesIO(content))
    ws = wb["Data"]
    assert ws["A1"].data_type == "s" and ws["A1"].value == "=SUM(1)"
    assert ws["A2"].data_type == "s" and ws["A2"].value == '=HYPERLINK("x")'
    assert ws["A3"].data_type == "s" and ws["A3"].value == "+381 11"
    assert ws["A4"].data_type == "s" and ws["A4"].value == "@cmd"
    assert ws["B2"].value == 1 and ws["B2"].data_type == "n"
    assert wb["Napomena"]["B1"].data_type == "s"
    assert wb["Napomena"]["B1"].value == "-2+3"


def test_xlsx_timeout_redirects_with_flag(monkeypatch):
    async def fake(db, **kwargs):
        return AnalitikaExportView(status="timeout")

    monkeypatch.setattr(routes_analitika.svc, "get_analitika_export", fake)
    resp = _client(sess=SESS).get("/app/analitika.xlsx?subject=awards&g1=winner")
    assert resp.status_code == 303
    assert resp.headers["location"] == (
        "/app/analitika?subject=awards&g1=winner&timeout=1"
    )


@pytest.mark.asyncio
async def test_timeout_flag_shows_hint_without_rerunning_report():
    db = _TimeoutSession()
    page = await svc.get_analitika_page(
        db, tid="t-1", params={"subject": "tenders", "timeout": "1"}, lang="sr"
    )
    assert page.timed_out is True
    assert db.rolled_back == 0
    assert not any("FROM tenders_normalized" in sql for sql in db.calls)


def _profile_row(**kw):
    base = dict(
        id="11111111-1111-1111-1111-111111111111",
        name="Šablon",
        classifications=[{"system_code": "eu_cpv", "code": "33000000"}],
        regions=[],
        keywords_include=[{"all_of": ["laptop"]}],
        keywords_exclude=["školu"],
        keyword_dict=False,
    )
    base.update(kw)
    return SimpleNamespace(**base)


def test_apply_profile_carries_keywords():
    spec = rb.parse_spec({"profile": "11111111-1111-1111-1111-111111111111"})
    out = svc._apply_profile(spec, [_profile_row()])
    assert out.profile_cpv == ("33",)
    assert out.profile_keywords == {
        "include": [{"all_of": ["laptop"]}],
        "exclude": ["školu"],
        "dict": False,
    }


def test_apply_profile_without_keywords_has_no_keyword_filter():
    spec = rb.parse_spec({"profile": "11111111-1111-1111-1111-111111111111"})
    out = svc._apply_profile(
        spec, [_profile_row(keywords_include=[], keywords_exclude=[])]
    )
    assert out.profile_keywords is None


@pytest.mark.asyncio
async def test_too_broad_keyword_scope_shows_timeout_hint_not_500(monkeypatch):
    async def boom(db, spec, **kw):
        raise rb.KeywordScopeTooBroad

    monkeypatch.setattr(svc.rb, "run_report", boom)

    class S(_TimeoutSession):
        async def execute(self, statement, params=None):
            if "FROM profiles" in str(statement):
                return SimpleNamespace(fetchall=lambda: [_profile_row()])
            return await super().execute(statement, params)

    page = await svc.get_analitika_page(
        S(),
        tid="t-1",
        params={
            "subject": "tenders",
            "profile": "11111111-1111-1111-1111-111111111111",
        },
        lang="sr",
    )
    assert page.status == "success" and page.timed_out is True


# --------------------------------------------------------------------------- #
#  Sorting by clicking the header (owner 10.10)
# --------------------------------------------------------------------------- #


def test_sort_parse_whitelist_and_defaults():
    p = rb.parse_spec
    assert p({"sort": "sum", "dir": "asc"}).sort == "sum"
    assert p({"sort": "sum", "dir": "asc"}).sort_dir == "asc"
    # default direction: numbers — descending, group — alphabetically
    assert p({"sort": "sum"}).sort_dir == "desc"
    assert p({"sort": "g1"}).sort_dir == "asc"
    assert p({"sort": "sum", "dir": "sideways"}).sort_dir == "desc"
    # unknown / foreign to the set / injection → no sorting (not 500)
    for bad in ("nope", "sum; DROP TABLE x", "n", "g2", "winners", "SUM "):
        assert p({"sort": bad, "dir": "asc"}).sort == ""
    assert p({"subject": "awards", "sort": "winners"}).sort == "winners"
    assert p({"g2": "region", "sort": "g2"}).sort == "g2"


def test_sort_sql_uses_whitelist_column_nulls_last_and_tiebreak():
    spec = rb.parse_spec(
        {"g1": "month", "g2": "region", "sort": "median", "dir": "asc"}
    )
    sql, params = rb.build_report_sql(spec, limit=50)
    assert (
        "ORDER BY median_eur ASC NULLS LAST, g1_key ASC NULLS LAST, "
        "g2_key ASC NULLS LAST\n    LIMIT" in sql
    )
    assert "median" not in params.values()
    # without an explicit sort the order stays as before
    sql0, _ = rb.build_report_sql(rb.parse_spec({"g1": "month"}), limit=50)
    assert "ORDER BY g1_key ASC NULLS LAST" in sql0


def test_sort_heads_toggle_indicator_and_keep_filters():
    spec = rb.parse_spec(
        {"subject": "tenders", "g1": "customer", "m": "sum", "region": "RS-00",
         "from": "2026-01-01", "sort": "avg", "dir": "asc"}
    )  # fmt: skip
    heads = svc.sort_heads(spec)
    assert set(heads) == {
        "g1", "count", "with_amount", "sum", "avg", "median", "customers"
    }  # fmt: skip
    cur = heads["avg"]
    assert (cur.aria_sort, cur.arrow) == ("ascending", "▲")
    assert "sort=avg" in cur.href and "dir=desc" in cur.href  # a repeated click
    assert "region=RS-00" in cur.href and "from=2026-01-01" in cur.href
    other = heads["count"]
    assert (other.aria_sort, other.arrow) == ("none", "")
    assert "sort=count" in other.href and "dir=desc" in other.href
    assert "dir=asc" in heads["g1"].href  # the first click on a group — alphabetically
    # without an explicit sort the indicator sits on the chosen metric (DESC)
    d = svc.sort_heads(rb.parse_spec({"subject": "tenders", "m": "sum"}))
    assert (d["sum"].aria_sort, d["sum"].arrow) == ("descending", "▼")
    assert "dir=asc" in d["sum"].href
    # chronology: by default ▲ on the group
    m = svc.sort_heads(rb.parse_spec({"subject": "tenders", "g1": "month"}))
    assert m["g1"].arrow == "▲"
    # winners — only on decisions
    assert "winners" in svc.sort_heads(rb.parse_spec({"subject": "awards"}))


def test_sort_export_href_carries_sort(monkeypatch):
    spec = rb.parse_spec({"subject": "tenders", "sort": "sum", "dir": "asc"})
    q = svc.spec_query(spec)
    assert (q["sort"], q["dir"]) == ("sum", "asc")
    assert "sort" not in svc.spec_query(rb.parse_spec({"subject": "tenders"}))


def test_page_renders_sortable_headers_with_aria_sort(monkeypatch):
    spec = rb.parse_spec(
        {"subject": "tenders", "g1": "month", "sort": "sum", "dir": "desc"}
    )
    _patch_page(
        monkeypatch,
        AnalitikaPageView(
            status="success",
            spec=spec,
            ran=True,
            rows=[RowView("2026-06", "", 2, 1, "100 €", "100 €", "100 €", 1, 0)],
            total=RowView("", "", 2, 1, "100 €", "100 €", "100 €", 1, 0),
            total_groups=1,
            heads=svc.sort_heads(spec),
        ),
    )
    html = _client(sess=SESS).get("/app/analitika?subject=tenders&lang=sr").text
    for key in ("g1", "count", "with_amount", "sum", "avg", "median", "customers"):
        assert f'data-testid="an-sort-{key}"' in html, key
    th_sum = re.search(r'<th[^>]*aria-sort="(\w+)"[^>]*><a[^>]*an-sort-sum', html)
    assert th_sum and th_sum.group(1) == "descending"
    th_cnt = re.search(r'<th[^>]*aria-sort="(\w+)"[^>]*><a[^>]*an-sort-count', html)
    assert th_cnt and th_cnt.group(1) == "none"
    assert "▼" in html and "dir=asc" in html  # a repeated click on “Zbir” → asc
    assert 'name="sort" value="sum"' in html  # the form does not lose the sorting
    assert "winners" not in html.split("<thead>")[1].split("</thead>")[0]


# --------------------------------------------------------------------------- #
#  The ready report: a result anchor, a loading indicator, the filters collapsed
# --------------------------------------------------------------------------- #


def test_preset_and_sort_links_anchor_to_result():
    presets = svc._presets([])
    assert presets and all(p.href.endswith("#an-result") for p in presets)
    heads = svc.sort_heads(rb.parse_spec({"subject": "tenders", "g1": "customer"}))
    assert heads and all(h.href.endswith("#an-result") for h in heads.values())


def test_export_link_has_no_anchor():
    # .xlsx is a file, not a page: an anchor in its address gives nothing.
    src = Path(svc.__file__).read_text(encoding="utf-8")
    export_line = next(ln for ln in src.splitlines() if "export_href=" in ln)
    assert "#an-result" not in export_line


@pytest.mark.parametrize(
    "view_kwargs",
    [
        {"ran": True, "rows": [RowView("Alfa", "", 1, 1, "1 €", "1 €", "1 €", 1, 0)]},
        {"ran": True, "timed_out": True},
        {"ran": True, "rows": []},
    ],
)
def test_every_result_outcome_is_anchor_target(monkeypatch, view_kwargs):
    spec = rb.parse_spec({"subject": "tenders"})
    _patch_page(
        monkeypatch,
        AnalitikaPageView(
            status="success",
            spec=spec,
            heads=svc.sort_heads(spec),
            presets=[],
            **view_kwargs,
        ),
    )
    html = _client(sess=SESS).get("/app/analitika?subject=tenders&lang=sr").text
    assert 'id="an-result" tabindex="-1"' in html


def test_form_submits_back_to_result_anchor(monkeypatch):
    _patch_page(
        monkeypatch,
        AnalitikaPageView(status="success", spec=rb.parse_spec({}), presets=[]),
    )
    html = _client(sess=SESS).get("/app/analitika?lang=sr").text
    assert 'action="/app/analitika#an-result"' in html


@pytest.mark.parametrize("ran,expect_open", [(False, True), (True, False)])
def test_filters_collapsed_once_report_ran(monkeypatch, ran, expect_open):
    _patch_page(
        monkeypatch,
        AnalitikaPageView(
            status="success", spec=rb.parse_spec({}), ran=ran, presets=[]
        ),
    )
    html = _client(sess=SESS).get("/app/analitika?lang=sr").text
    m = re.search(r'<details class="pf__more"( open)?>', html)
    assert m and bool(m.group(1)) is expect_open


@pytest.mark.parametrize(
    "lang,label",
    [("sr", "Učitavanje…"), ("ru", "Загрузка…"), ("en", "Loading…")],
)
def test_busy_label_and_hooks_rendered(monkeypatch, lang, label):
    spec = rb.parse_spec({})
    preset = svc.Preset(
        key="an_preset_monthly", href="/app/analitika?g1=month#an-result"
    )
    _patch_page(
        monkeypatch,
        AnalitikaPageView(status="success", spec=spec, presets=[preset]),
    )
    html = _client(sess=SESS).get(f"/app/analitika?lang={lang}").text
    assert f"data-an-busy hidden>{label}" in html
    assert "data-an-preset" in html
    assert "data-an-run" in html


def test_busy_script_wired_with_sri():
    root = Path(__file__).resolve().parents[2]
    digest = hashlib.sha384(
        (root / "src/api/static/js/analitika_busy.js").read_bytes()
    ).digest()
    expected = "sha384-" + base64.b64encode(digest).decode()
    html = (root / "src/api/templates/analitika.html").read_text(encoding="utf-8")
    m = re.search(
        r"static_url\('js/analitika_busy\.js'\) \}\}\" integrity=\"([^\"]+)\"", html
    )
    assert m and m.group(1) == expected


# --------------------------------------------------------------------------- #
#  The active preset: the highlight matches the parameters of the current report
# --------------------------------------------------------------------------- #

_WINNERS_HREF = 'href="/app/analitika?subject=awards&amp;g1=winner&amp;m=sum#an-result"'


async def _none(*_a, **_k):
    return None


def _render_real_page(monkeypatch, query: str, profiles: list | None = None) -> str:
    """The real ``get_analitika_page``: only the gate, the profiles and the selection are mocked."""
    monkeypatch.setattr(svc, "_gate", _none)
    monkeypatch.setattr(svc, "_run_guarded", _none)

    async def load_profiles(db, tid):
        return list(profiles or [])

    monkeypatch.setattr(svc, "_load_profiles", load_profiles)
    resp = _client(sess=SESS).get(f"/app/analitika?{query}&lang=sr")
    assert resp.status_code == 200
    return resp.text


def _preset_anchors(html: str) -> list[str]:
    block = html.split('data-testid="an-presets"', 1)[1].split("</section>", 1)[0]
    return re.findall(r"<a [^>]*data-an-preset[^>]*>", block)


def test_winners_preset_marked_active(monkeypatch):
    html = _render_real_page(monkeypatch, "subject=awards&g1=winner&m=sum")
    active = [a for a in _preset_anchors(html) if 'aria-current="true"' in a]
    assert len(active) == 1
    assert _WINNERS_HREF in active[0]
    assert "is-active" in active[0]
    assert all(
        "aria-current" not in a for a in _preset_anchors(html) if a not in active
    )


def test_active_preset_survives_sort_and_page(monkeypatch):
    html = _render_real_page(
        monkeypatch, "subject=awards&g1=winner&m=sum&sort=g1&dir=asc&page=2"
    )
    active = [a for a in _preset_anchors(html) if 'aria-current="true"' in a]
    assert len(active) == 1 and _WINNERS_HREF in active[0]


def test_changed_filter_leaves_no_preset_active(monkeypatch):
    html = _render_real_page(monkeypatch, "subject=awards&g1=customer&m=sum")
    assert not [a for a in _preset_anchors(html) if "aria-current" in a]
    assert "is-active" not in " ".join(_preset_anchors(html))


_PROFILE_ID = "6f1d2c3a-0000-4000-8000-000000000001"


def _preset_profile_row():
    return SimpleNamespace(
        id=_PROFILE_ID,
        name="Мой шаблон",
        classifications=[],
        regions=[],
        keywords_include=[],
        keywords_exclude=[],
        keyword_dict=False,
    )


def test_region_filter_leaves_no_preset_active(monkeypatch):
    html = _render_real_page(monkeypatch, "subject=awards&g1=winner&m=sum&region=RS-00")
    assert not [a for a in _preset_anchors(html) if "aria-current" in a]
    assert "is-active" not in " ".join(_preset_anchors(html))


def test_preset_with_profile_active_until_date_filter_added(monkeypatch):
    base = f"subject=awards&g1=winner&m=sum&profile={_PROFILE_ID}"
    html = _render_real_page(monkeypatch, base, profiles=[_preset_profile_row()])
    active = [a for a in _preset_anchors(html) if 'aria-current="true"' in a]
    assert len(active) == 1
    assert "g1=winner" in active[0] and f"profile={_PROFILE_ID}" in active[0]
    html = _render_real_page(
        monkeypatch, base + "&from=2026-01-01", profiles=[_preset_profile_row()]
    )
    assert not [a for a in _preset_anchors(html) if "aria-current" in a]
