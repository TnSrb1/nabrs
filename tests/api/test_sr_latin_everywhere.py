"""Guard (owner, 10.10.2026): “Serbian EVERYWHERE in Latin script only”.

With ``lang == “sr”`` the page contains NOT A SINGLE Cyrillic letter, even when
the data (title, customer, CPV reference, districts) is stored in the DB in Cyrillic.
The data is not changed — transliteration happens at the display boundary
(:class:`src.api.templating._SrLatinTemplate`). ru/en — unchanged.
The expectation is a literal (``[U+0400-U+04FF]`` = 0), not the same ``cyr_to_lat``.
"""

from __future__ import annotations

import re
import time
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient

from src.api import deps
from src.api.main import app
from src.api.public import geo_cache
from tests.api.conftest import FakeResult, FakeSession

CYR = re.compile("[Ѐ-ӿ]")
_TID = "11111111-1111-4111-8111-111111111111"


def _client(session: FakeSession) -> TestClient:
    def factory():
        return session

    app.dependency_overrides[deps.get_session_factory] = lambda: factory
    app.dependency_overrides[deps.get_optional_session] = lambda: None
    return TestClient(app)


def teardown_function():
    app.dependency_overrides.clear()


def _row(**kw):
    base = dict(
        id="t-1",
        title="Радови на реконструкцији пута",
        customer_name="ЈП Путеви Србије",
        deadline_at=datetime.now(timezone.utc) + timedelta(days=10),
        amount_eur=50000.0,
    )
    base.update(kw)
    return SimpleNamespace(**base)


def _tender(**kw):
    base = dict(
        id=_TID,
        title="Видеонаблюдение за школу — Радови",
        customer_name="ОШ «Вук Караџић»",
        amount_eur=12345.0,
        is_amount_estimated=False,
        deadline_at=datetime.now(timezone.utc) + timedelta(days=20),
        status="active",
        regions=["RS-00"],
        procedure_type=None,
        description="Опис набавке: израда пројекта.",
        classifications=None,
        source_url="https://jnportal.rs/t/1",
        canonical_tender_id=None,
    )
    base.update(kw)
    return SimpleNamespace(**base)


def _prime():
    geo_cache._cpv_cache = (
        time.monotonic() + 3600,
        {
            "45": {
                "total": 10,
                "open": 2,
                "open_sum_eur": 75000.0,
                "label": "Грађевински радови",
            }
        },
    )
    geo_cache._region_cache = (
        time.monotonic() + 3600,
        {"RS-00": {"total": 20, "open": 1, "open_sum_eur": None}},
    )


def _cpv_page(lang):
    _prime()
    s = FakeSession(results=[FakeResult(items=[_row()]), FakeResult(items=[_row()])])
    return _client(s).get(f"/tenderi/cpv/45?lang={lang}")


def _region_page(lang):
    _prime()
    s = FakeSession(results=[FakeResult(items=[_row()]), FakeResult(items=[])])
    return _client(s).get(f"/tenderi/region/RS-00?lang={lang}")


def _hub_page(lang):
    _prime()
    return _client(FakeSession(results=[])).get(f"/tenderi?lang={lang}")


def _tender_page(lang):
    s = FakeSession(results=[FakeResult(items=[_tender()])])
    return _client(s).get(f"/tender/{_TID}?lang={lang}")


PAGES = {
    "tenderi_cpv": _cpv_page,
    "tenderi_region": _region_page,
    "tenderi_hub": _hub_page,
    "tender_card": _tender_page,
}


@pytest.mark.parametrize("name", sorted(PAGES))
def test_sr_page_has_zero_cyrillic(name):
    resp = PAGES[name]("sr")
    assert resp.status_code == 200
    assert CYR.findall(resp.text) == []


def test_sr_page_shows_transliterated_data():
    """The data are not lost, but transliterated (a literal from the fixture)."""
    text = _cpv_page("sr").text
    assert "Radovi na rekonstrukciji puta" in text
    assert "JP Putevi Srbije" in text


def test_ru_page_keeps_cyrillic():
    """ru sees the Cyrillic as it is — the mechanism does not touch the other languages."""
    assert "Радови на реконструкцији пута" in _cpv_page("ru").text


# --------------------------------------------------------------------------- #
#  Outside Jinja: CSV, Excel, e-mails, Telegram, llms.txt
# --------------------------------------------------------------------------- #

_CYR_TITLE = "Радови на путу — Ђорђе Љубиша ЈП"


def test_csv_sr_zero_cyrillic_and_other_langs_untouched():
    from src.bots.analytics_export import csv_bytes

    rows = [["Назив", _CYR_TITLE]]
    assert CYR.findall(csv_bytes(rows, "sr").decode("utf-8")) == []
    assert "Радови" in csv_bytes(rows, "ru").decode("utf-8")
    assert "Радови" in csv_bytes(rows).decode("utf-8")


def test_my_tenders_csv_sr_zero_cyrillic():
    from src.api.cabinet.my_tenders_export import build_my_tenders_csv

    row = SimpleNamespace(
        title=_CYR_TITLE,
        customer_name="ЈП Путеви Србије",
        amount=None,
        currency="RSD",
        deadline_at=None,
        user_status=None,
        folder_names="Папка",
        tender_url="https://nab.rs/tender/1",
    )
    out = build_my_tenders_csv([row], "sr").decode("utf-8")  # type: ignore[list-item]
    assert CYR.findall(out) == []
    assert "Djordje" in out or "Đorđe" in out


def test_xlsx_sr_zero_cyrillic():
    import io

    from openpyxl import load_workbook

    from src.api.cabinet.analitika_xlsx import build_xlsx

    raw = build_xlsx(["Назив"], [[_CYR_TITLE]], [("Напомена", "Тест")], "sr")
    wb = load_workbook(io.BytesIO(raw))
    cells = [str(c.value) for ws in wb for row in ws.iter_rows() for c in row]
    assert CYR.findall(" ".join(cells)) == []
    assert "Radovi na putu" in " ".join(cells)


@pytest.mark.asyncio
async def test_email_sr_zero_cyrillic_ru_untouched():
    from src.api.emailer import MockEmailSender

    sender = MockEmailSender()
    await sender.send(
        to="a@b.rs",
        subject=_CYR_TITLE,
        text_body=_CYR_TITLE,
        html_body=f"<p>{_CYR_TITLE}</p>",
        locale="sr",
    )
    await sender.send(to="a@b.rs", subject=_CYR_TITLE, text_body="x", locale="ru")
    sr, ru = sender.outbox
    assert CYR.findall(f"{sr.subject}{sr.text_body}{sr.html_body}") == []
    assert ru.subject == _CYR_TITLE


def test_telegram_method_sr_zero_cyrillic_ru_untouched():
    from aiogram.methods import SendMessage
    from aiogram.types import InlineKeyboardButton, InlineKeyboardMarkup

    from src.bots.sr_latin import latinize_method

    kb = InlineKeyboardMarkup(
        inline_keyboard=[[InlineKeyboardButton(text="Учествујем", callback_data="x")]]
    )
    m = SendMessage(chat_id=1, text=_CYR_TITLE, reply_markup=kb)
    out = latinize_method(m, "sr")
    assert CYR.findall(out.text + out.reply_markup.inline_keyboard[0][0].text) == []
    assert out.reply_markup.inline_keyboard[0][0].callback_data == "x"
    assert latinize_method(m, "ru").text == _CYR_TITLE


def test_llms_txt_sr_zero_cyrillic():
    from src.api.public import routes_seo

    s = FakeSession(results=[])
    resp = _client(s).get("/llms.txt")
    assert resp.status_code == 200
    assert CYR.findall(resp.text) == []
    assert routes_seo.sr_display("Ћирилица", "sr") == "Ćirilica"


# --------------------------------------------------------------------------- #
#  The other render points: /statistika (CPV names, customers), the static
#  of the showcase, the tender card in the cabinet
# --------------------------------------------------------------------------- #


def test_statistika_sr_zero_cyrillic(monkeypatch):
    from dataclasses import replace

    from src.api.public import stats_page
    from tests.api.test_statistika_page import STATS

    stats = replace(
        STATS,
        top_customers=(
            stats_page.CustomerRow("Град Београд", 3, STATS.awards_sum_eur),
        ),
    )

    async def fake_stats(_factory):
        return stats

    monkeypatch.setattr(stats_page, "get_public_stats", fake_stats)
    app.dependency_overrides[deps.get_session_factory] = lambda: FakeSession
    app.dependency_overrides[deps.get_optional_session] = lambda: None
    try:
        resp = TestClient(app).get("/statistika?lang=sr")
    finally:
        stats_page.clear_cache()
    assert resp.status_code == 200
    assert "Grad Beograd" in resp.text
    assert CYR.findall(resp.text) == []


@pytest.mark.parametrize("path", ["/", "/help", "/plans", "/izvori", "/kontakt"])
def test_static_public_pages_sr_zero_cyrillic(path):
    s = FakeSession(results=[])
    resp = _client(s).get(f"{path}?lang=sr")
    assert resp.status_code == 200
    # Without exceptions (owner 10.10): the Serbian interface EVERYWHERE is only in Latin,
    # including the example of the AI analysis on the home page (the translation to sr is in the template).
    assert CYR.findall(resp.text) == []


def test_cabinet_tender_detail_sr_zero_cyrillic(monkeypatch):
    from tests.api import test_tender_description_markers as m

    monkeypatch.setattr(
        m.routes_tender_ai.svc,
        "get_tender_analyses",
        m._async(m._cabinet_data("### Документ.pdf\nОпис: израда пројекта.")),
    )
    resp = m._cabinet_client().get(
        f"/app/tenders/{m._TENDER}?profile_id={m._PROFILE}&lang=sr"
    )
    assert resp.status_code == 200
    assert "Nabavka računara" in resp.text
    assert CYR.findall(resp.text) == []


def test_every_template_response_goes_through_sr_latin_class():
    """One mechanism: the Jinja environment renders with a class that Latinises sr."""
    from src.api.templating import _SrLatinTemplate, templates

    assert templates.env.template_class is _SrLatinTemplate


# --- PR #647 review: form values going to the server stay as in the DB ----


def _confirm_html(lang):
    from types import SimpleNamespace as NS

    from starlette.datastructures import QueryParams

    from src.api.cabinet.i18n import cabinet_context_for_lang
    from src.api.templating import templates

    ctx = dict(
        cabinet_context_for_lang(lang),
        request=NS(
            url=NS(path="/x"), state=NS(), query_params=QueryParams(""), cookies={}
        ),
        profile={"name": "Профил", "keywords_include": [{"all_of": ["щебень"]}]},
        count=3,
        keywords_include_text="щебень\nпесак",
        keywords_exclude_text="",
        cls_include_text="Камен",
        cls_exclude_text="",
        schedule_choices=[],
        source_choices=[],
        prefill_source=None,
        edit_query="",
    )
    return templates.get_template("profile_confirm.html").render(**ctx)


def _cyr(s):
    return re.findall(r"[Ѐ-ӿ]", s)


def test_sr_profile_confirm_keeps_textarea_values_cyrillic():
    out = _confirm_html("sr")
    assert (
        '<textarea name="keywords_include" class="hidden">щебень\nпесак</textarea>'
        in out
    )
    assert '<textarea name="cls_include" class="hidden">Камен</textarea>' in out
    # The visible — in Latin. ONLY the values of the form fields are excluded: the content of
    # textarea and value= of input (explicitly and narrowly, not the whole tag).
    visible = re.sub(r"<textarea\b.*?</textarea>", "", out, flags=re.S)
    visible = re.sub(r'(\svalue=")[^"]*"', r'\1"', visible)
    assert _cyr(visible) == []
    assert 'name="name" value="Профил"' in out
    assert "Profil" in visible


def test_sr_html_keeps_option_value_and_data_attrs_but_translits_text():
    from src.api.templating import _sr_latin_html

    html = (
        '<select><option value="Нови Сад">Нови Сад</option></select>'
        '<button value="щ" data-code="Ђ" data-confirm="Да ли сте сигурни?">Да</button>'
        '<textarea name="n">Бележка</textarea>'
    )
    out = _sr_latin_html(html)
    assert '<option value="Нови Сад">Novi Sad</option>' in out
    assert 'value="щ"' in out and 'data-code="Ђ"' in out
    assert 'data-confirm="Da li ste sigurni?"' in out
    assert ">Bilješka<" not in out and ">Бележка<" in out
    assert ">Da</button>" in out


def test_ru_profile_confirm_unchanged():
    assert "щебень" in _confirm_html("ru")
