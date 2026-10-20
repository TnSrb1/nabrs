"""Integration: the sorting of the «Analitika» table by each column (owner 10.10).

A live PG under the application role. The suffix ``sfx`` in the customer names isolates
the rows from the rest of the DB; each report is filtered by ``customer=sfx``.

Tenders (grouped by the publication month), the expectations are counted by hand:

    month    tender  customer  EUR      →  n  with sum  sum    average  median  customers
    2026-03  t1      c1        1000        1  1         1000   1000     1000     1
    2026-04  t2      c1        100         3  2         300    150      150      3
             t3      c2        200
             t4      c3        NULL
    2026-05  t5      c1        NULL        2  0         NULL   NULL     NULL     1
             t6      c1        NULL
    2026-06  t7      c1        50          2  2         500    250      250      2
             t8      c2        450

Decisions (the month of the decision; the winner):

    month    decision  winner  EUR   →  winners
    2026-03  a1        w1      10       1
    2026-04  a2        w1      20       2
             a3        w2      NULL
    2026-05  a4        w3      5        1
"""

from __future__ import annotations

import io
import json
import os
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlalchemy.pool import NullPool

from src.analytics import report_builder as rb
from src.api.cabinet import analitika_service as svc

TEST_URL = os.getenv("TEST_RLS_DATABASE_URL", "")
SUPER_URL = os.getenv("TEST_SUPERUSER_DATABASE_URL", TEST_URL)

pytestmark = pytest.mark.skipif(
    not TEST_URL,
    reason="TEST_RLS_DATABASE_URL не задан — интеграция сортировки пропускается",
)

M3, M4, M5, M6 = "2026-03", "2026-04", "2026-05", "2026-06"

# (month, customer, EUR)
_TENDERS = [
    ("2026-03-10", "c1", Decimal("1000")),
    ("2026-04-10", "c1", Decimal("100")),
    ("2026-04-11", "c2", Decimal("200")),
    ("2026-04-12", "c3", None),
    ("2026-05-10", "c1", None),
    ("2026-05-11", "c1", None),
    ("2026-06-10", "c1", Decimal("50")),
    ("2026-06-11", "c2", Decimal("450")),
]
# (uid, decision date, winner, EUR)
_AWARDS = [
    ("a1", "2026-03-15", "w1", Decimal("10")),
    ("a2", "2026-04-15", "w1", Decimal("20")),
    ("a3", "2026-04-16", "w2", None),
    ("a4", "2026-05-15", "w3", Decimal("5")),
]


def _ts(d: str) -> datetime:
    return datetime.fromisoformat(d + "T10:00:00+00:00")


@pytest_asyncio.fixture
async def world():
    eng = create_async_engine(SUPER_URL, poolclass=NullPool, future=True)
    sfx = uuid.uuid4().hex[:10]
    ids: list[uuid.UUID] = []
    async with eng.begin() as c:
        for i, (pub, cust, amt) in enumerate(_TENDERS):
            tid = uuid.uuid4()
            ids.append(tid)
            await c.execute(
                text(
                    "INSERT INTO tenders_normalized (id, source_key, source_uid, "
                    "country_code, status, customer_name, published_at, "
                    "first_seen_at, amount_eur, classifications, "
                    "customer_region_code, regions, procedure_type) "
                    "VALUES (:id, 'jnportal', :uid, 'RS', 'active', :cust, :pub, "
                    ":pub, :amt, CAST(:cls AS jsonb), 'RS-00', ARRAY['RS-00'], 'open')"
                ),
                {
                    "id": tid,
                    "uid": f"srt-{sfx}-{i}",
                    "cust": f"Naručilac {cust} {sfx}",
                    "pub": _ts(pub),
                    "amt": amt,
                    "cls": json.dumps([{"scheme": "CPV", "code": "45210000"}]),
                },
            )
        for uid, dd, w, amt in _AWARDS:
            await c.execute(
                text(
                    "INSERT INTO tender_awards (tender_id, source_uid, "
                    "document_type_id, customer_name, winner_pib, winner_name, "
                    "decision_date, award_amount_eur, raw) VALUES (:t, :uid, 11, "
                    ":cust, :wp, :wn, CAST(:d AS date), :amt, CAST(:raw AS jsonb))"
                ),
                {
                    "t": ids[0],
                    "uid": f"srt-{sfx}-{uid}",
                    "cust": f"Naručilac c1 {sfx}",
                    "wp": f"{w}-{sfx}",
                    "wn": f"Firma {w}",
                    "d": date.fromisoformat(dd),
                    "amt": amt,
                    "raw": json.dumps({"contract_row": {"IsReversed": False}}),
                },
            )
    app_eng = create_async_engine(TEST_URL, poolclass=NullPool, future=True)
    try:
        yield sfx, app_eng
    finally:
        async with eng.begin() as c:
            await c.execute(
                text("DELETE FROM tenders_normalized WHERE id = ANY(:ids)"),
                {"ids": ids},
            )
        await app_eng.dispose()
        await eng.dispose()


async def _run(app_eng, *, limit=50, **params):
    spec = rb.parse_spec(params)
    async with AsyncSession(app_eng) as s:
        return spec, await rb.run_report(s, spec, limit=limit, with_coverage=True)


async def _keys(app_eng, **params) -> list[str]:
    _, res = await _run(app_eng, **params)
    return [r.g1_key for r in res.rows]


# (key, asc, desc) — literals from the table in the docstring; the tie-break is always g1_key ASC.
_EXPECT = [
    ("g1", [M3, M4, M5, M6], [M6, M5, M4, M3]),
    ("count", [M3, M5, M6, M4], [M4, M5, M6, M3]),
    ("with_amount", [M5, M3, M4, M6], [M4, M6, M3, M5]),
    ("sum", [M4, M6, M3, M5], [M3, M6, M4, M5]),  # M5 (NULL) — at the end in both
    ("avg", [M4, M6, M3, M5], [M3, M6, M4, M5]),
    ("median", [M4, M6, M3, M5], [M3, M6, M4, M5]),
    ("customers", [M3, M5, M6, M4], [M4, M6, M3, M5]),
]


@pytest.mark.asyncio
@pytest.mark.parametrize("key,asc,desc", _EXPECT)
async def test_tenders_every_column_asc_and_desc(world, key, asc, desc):
    sfx, app_eng = world
    base = dict(subject="tenders", g1="month", customer=sfx, sort=key)
    assert await _keys(app_eng, **base, dir="asc") == asc
    assert await _keys(app_eng, **base, dir="desc") == desc


@pytest.mark.asyncio
async def test_null_amounts_last_in_both_directions(world):
    sfx, app_eng = world
    for key in ("sum", "avg", "median"):
        for d in ("asc", "desc"):
            keys = await _keys(
                app_eng, subject="tenders", g1="month", customer=sfx, sort=key, dir=d
            )
            assert keys[-1] == M5, (key, d)


@pytest.mark.asyncio
async def test_limit_applies_after_sort(world):
    sfx, app_eng = world
    # By default (chronology) the first two are March/April; by sum asc — April/June.
    assert await _keys(
        app_eng, limit=2, subject="tenders", g1="month", customer=sfx
    ) == [M3, M4]
    assert await _keys(
        app_eng,
        limit=2,
        subject="tenders",
        g1="month",
        customer=sfx,
        sort="sum",
        dir="asc",
    ) == [M4, M6]
    # total_groups — up to the limit.
    _, res = await _run(
        app_eng, limit=2, subject="tenders", g1="month", customer=sfx, sort="sum"
    )
    assert res.total_groups == 4


@pytest.mark.asyncio
async def test_unknown_sort_key_falls_back_to_default_order(world):
    sfx, app_eng = world
    default = await _keys(app_eng, subject="tenders", g1="month", customer=sfx)
    assert default == [M3, M4, M5, M6]
    for bad in ("nope", "n; DROP TABLE tenders_normalized", "g2", "winners", ""):
        assert (
            await _keys(
                app_eng,
                subject="tenders",
                g1="month",
                customer=sfx,
                sort=bad,
                dir="desc",
            )
            == default
        )


@pytest.mark.asyncio
async def test_group2_sort_with_deterministic_tiebreak(world):
    sfx, app_eng = world
    _, res = await _run(
        app_eng,
        subject="tenders",
        g1="month",
        g2="customer",
        customer=sfx,
        sort="g2",
        dir="desc",
    )
    pairs = [(r.g1_key, r.g2_key) for r in res.rows]
    assert [p[1].split()[1] for p in pairs] == [
        "c3",
        "c2",
        "c2",
        "c1",
        "c1",
        "c1",
        "c1",
    ]
    # within the same customer — the months in ascending order (the tie-break by g1_key)
    assert [p[0] for p in pairs if p[1].split()[1] == "c1"] == [M3, M4, M5, M6]


@pytest.mark.asyncio
async def test_customer_group_sorts_by_label(world):
    sfx, app_eng = world
    asc = await _keys(
        app_eng, subject="tenders", g1="customer", customer=sfx, sort="g1", dir="asc"
    )
    assert [k.split()[1] for k in asc] == ["c1", "c2", "c3"]
    desc = await _keys(
        app_eng, subject="tenders", g1="customer", customer=sfx, sort="g1", dir="desc"
    )
    assert [k.split()[1] for k in desc] == ["c3", "c2", "c1"]


@pytest.mark.asyncio
async def test_awards_sort_including_winners_column(world):
    sfx, app_eng = world
    base = dict(subject="awards", g1="month", customer=sfx)
    # decisions: M3 1 winner, M4 2, M5 1
    assert await _keys(app_eng, **base, sort="winners", dir="asc") == [M3, M5, M4]
    assert await _keys(app_eng, **base, sort="winners", dir="desc") == [M4, M3, M5]
    # sums: M3 10, M4 20, M5 5
    assert await _keys(app_eng, **base, sort="sum", dir="asc") == [M5, M3, M4]
    assert await _keys(app_eng, **base, sort="sum", dir="desc") == [M4, M3, M5]
    # g1 for the winner: the labels of the winners, both ways
    w_asc = await _keys(
        app_eng, subject="awards", g1="winner", customer=sfx, sort="g1", dir="asc"
    )
    assert [k.split("-")[0] for k in w_asc] == ["w1", "w2", "w3"]


@pytest.mark.asyncio
async def test_excel_has_same_order_as_screen(world):
    sfx, app_eng = world
    spec, res = await _run(
        app_eng, subject="tenders", g1="month", customer=sfx, sort="sum", dir="desc"
    )
    screen = [r.g1_key for r in res.rows]
    assert screen == [M3, M6, M4, M5]
    blob = svc.build_export(
        spec, res, lang="sr", generated_at=datetime(2026, 10, 10, tzinfo=timezone.utc)
    )
    ws = load_workbook(io.BytesIO(blob))["Data"]
    col = [row[0].value for row in ws.iter_rows(min_row=2)][:-1]  # without «Ukupno»
    assert [v.strftime("%Y-%m") for v in col] == screen
